import sqlite3

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.carga import Carga, ErrorCarga
from src.extraccion import Extraccion
from src.transformacion import Transformacion


@pytest.fixture
def tablas(cliente_mongo):
    with Extraccion(base_datos="airbnb_test", cliente=cliente_mongo) as ext:
        return Transformacion().ejecutar(ext.extraer_todo())


def test_carga_sqlite_y_verificacion(tablas, tmp_path):
    carga = Carga(tmp_path / "t.db", tmp_path / "xlsx")
    carga.cargar_sqlite(tablas)
    reporte = carga.verificar_carga(tablas)
    assert reporte["ok"].all()
    with sqlite3.connect(tmp_path / "t.db") as con:
        n = con.execute("SELECT COUNT(*) FROM listings WHERE rango_precio = 'Sin precio'").fetchone()[0]
        idx = {f[1] for f in con.execute("SELECT * FROM sqlite_master WHERE type='index'")}
    assert n == 1 and "ux_listings" in idx and "ux_reviews" in idx


def test_la_carga_es_idempotente(tablas, tmp_path):
    carga = Carga(tmp_path / "t.db", tmp_path / "xlsx")
    carga.cargar_sqlite(tablas)
    carga.cargar_sqlite(tablas)                       # segunda ejecución: no duplica
    assert carga.verificar_carga(tablas)["ok"].all()


def test_verificacion_detecta_diferencias(tablas, tmp_path):
    carga = Carga(tmp_path / "t.db", tmp_path / "xlsx")
    carga.cargar_sqlite(tablas)
    with sqlite3.connect(tmp_path / "t.db") as con:
        con.execute("DELETE FROM reviews")
    with pytest.raises(ErrorCarga, match="diferencias"):
        carga.verificar_carga(tablas)


def test_un_duplicado_no_resuelto_rompe_el_indice_unico(tablas, tmp_path):
    malas = dict(tablas)
    malas["listings"] = pd.concat([tablas["listings"], tablas["listings"].iloc[[0]]], ignore_index=True)
    with pytest.raises(ErrorCarga):
        Carga(tmp_path / "t.db", tmp_path / "xlsx").cargar_sqlite(malas)


def test_exporta_xlsx_sin_comentarios(tablas, tmp_path):
    archivos = Carga(tmp_path / "t.db", tmp_path / "xlsx").exportar_xlsx(tablas)
    nombres = {a.name for a in archivos}
    assert {"airbnb_listings.xlsx", "airbnb_reviews.xlsx", "airbnb_calendario_mensual.xlsx"} <= nombres
    hoja = load_workbook(tmp_path / "xlsx" / "airbnb_reviews.xlsx")["reviews"]
    encabezados = [c.value for c in hoja[1]]
    assert "comentario" not in encabezados and "review_id" in encabezados and hoja.max_row == 4
    libro = load_workbook(tmp_path / "xlsx" / "airbnb_listings.xlsx")
    assert set(libro.sheetnames) == {"listings", "hosts"}


def test_ids_largos_se_exportan_como_texto_en_xlsx(tmp_path):
    # Excel guarda números como double (15 dígitos): 1234567890123456789 y ...790 colapsarían en uno.
    ids = [1234567890123456789, 1234567890123456790]
    tablas = {"listings": pd.DataFrame({"listing_id": ids, "host_id": [10.0, None], "price": [1.5, 2.5]})}
    Carga(tmp_path / "t.db", tmp_path / "xlsx").exportar_xlsx(tablas, plan={"l.xlsx": ["listings"]})
    hoja = load_workbook(tmp_path / "xlsx" / "l.xlsx")["listings"]
    assert [hoja.cell(r, 1).value for r in (2, 3)] == [str(i) for i in ids]
    assert hoja.cell(2, 2).value == "10" and hoja.cell(3, 2).value is None
    assert hoja.cell(2, 3).value == 1.5                      # las columnas que no son ID siguen numéricas
